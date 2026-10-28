# -*- coding: utf-8 -*-
"""workday.py — 工作日历工具（通用件，不含任何业务核对规则）

读一份「日期 / 是否工作日」两列的 xlsx 日历，回答两个问题：
  1. 今天是不是工作日（日历不可用时按周一至周五兜底）
  2. 上一个工作日是哪天

中国节假日与调休都体现在日历的这一列：节假日标「否」，调休上班的周末标「是」。

CLI（与调度脚本的约定一致）：
  python workday.py --is-workday   --calendar cal.xlsx
      -> CALENDAR_OK=1/0 + WORKDAY=1/0
  python workday.py --prev-workday --calendar cal.xlsx
      -> PREV_WORKDAY=YYYY-MM-DD（读不到日历时按周一至周五倒推）
"""
import argparse
import os
import sys
from datetime import date, datetime, timedelta

from openpyxl import load_workbook

# 日历里代表「上班」的写法，大小写不敏感
TRUE_FLAGS = ("Y", "YES", "TRUE", "1", "是", "工作日")


def norm(v):
    """把任意单元格值规整成用于表头匹配的字符串（去空白 + 小写）"""
    return str(v).strip().lower() if v is not None else ""


def to_date(v):
    """把单元格值转成 date：支持 datetime / date / Excel 序列号 / 常见字符串日期"""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):
        # Excel 序列号：以 1899-12-30 为原点
        try:
            return (datetime(1899, 12, 30) + timedelta(days=float(v))).date()
        except Exception:
            return None
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%d.%m.%Y", "%d/%m/%Y", "%Y%m%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except Exception:
            continue
    return None


def find_col(header, names, prefix=None):
    """在表头行中按名称找列号（1 起）。names 精确匹配，prefix 作前缀兜底。"""
    for want in names:
        w = norm(want)
        for i, v in enumerate(header):
            if norm(v) == w:
                return i + 1
    if prefix:
        p = norm(prefix)
        for i, v in enumerate(header):
            if norm(v).startswith(p):
                return i + 1
    return None


def load_workdays(cal_path):
    """从工作日历读出所有标记为工作日的日期集合

    日历不可用或表头识别失败时返回 None，由调用方决定如何兜底。
    """
    if not cal_path or not os.path.isfile(cal_path):
        return None
    wb = load_workbook(cal_path, data_only=True)
    try:
        # 优先选名字里含「工作日」的表，否则用第一个表
        ws = next((wb[s] for s in wb.sheetnames if "工作日" in s), wb[wb.sheetnames[0]])
        c_date = c_wd = None
        hdr = 0
        # 在前 10 行内找表头行（同时含「日期」和「是否工作日」）
        for r in range(1, min(ws.max_row, 10) + 1):
            vals = [ws.cell(row=r, column=c).value for c in range(1, min(ws.max_column, 10) + 1)]
            d = find_col(vals, ["日期"], prefix="日期")
            w = find_col(vals, ["是否工作日", "工作日", "是否上班"], prefix="是否工作日")
            if d and w:
                c_date, c_wd, hdr = d, w, r
                break
        if not c_date:
            return None
        days = set()
        for r in range(hdr + 1, ws.max_row + 1):
            dv = to_date(ws.cell(row=r, column=c_date).value)
            if dv is None:
                continue
            wv = ws.cell(row=r, column=c_wd).value
            flag = str(wv).strip().upper() if wv is not None else ""
            if flag in TRUE_FLAGS:
                days.add(dv)
        return days
    finally:
        wb.close()


def load_prev_workday(cal_path, today):
    """从工作日历读出「上一个工作日」（严格早于 today 的最近一个工作日）"""
    days = load_workdays(cal_path)
    if not days:
        return None
    prev = [d for d in days if d < today]
    return max(prev) if prev else None


def is_workday(cal_path, today):
    """判断 today 是否工作日，返回 (是否工作日, 日历是否可用)

    日历不可用时按周一至周五兜底：宁可多跑一次，也不要让定时任务因为
    日历文件缺失而静默停摆（节假日/调休仍以日历为准）。
    注意区分「读不到日历」（None）与「读到了但没有任何工作日」（空集合）：
    后者说明日历有效、只是当天不是工作日，不能走兜底。
    """
    days = load_workdays(cal_path)
    if days is None:
        return (today.weekday() < 5, False)
    return (today in days, True)


def prev_workday_by_weekday(today):
    """日历不可用时的兜底：按周一至周五倒推上一个工作日（不识别法定节假日）"""
    d = today - timedelta(days=1)
    while d.weekday() >= 5:
        d -= timedelta(days=1)
    return d


def main(argv=None):
    p = argparse.ArgumentParser(description="工作日历判断（通用）")
    p.add_argument("--calendar", default="", help="工作日历 xlsx")
    p.add_argument("--is-workday", action="store_true", help="打印 CALENDAR_OK / WORKDAY")
    p.add_argument("--prev-workday", action="store_true", help="打印 PREV_WORKDAY")
    args = p.parse_args(argv)
    today = date.today()
    if args.is_workday:
        ok, has_cal = is_workday(args.calendar, today)
        print("CALENDAR_OK=%d" % (1 if has_cal else 0))
        print("WORKDAY=%d" % (1 if ok else 0))
        return 0
    if args.prev_workday:
        d = load_prev_workday(args.calendar, today) or prev_workday_by_weekday(today)
        print("PREV_WORKDAY=%s" % d.isoformat())
        return 0
    p.print_help()
    return 2


if __name__ == "__main__":
    sys.exit(main())
